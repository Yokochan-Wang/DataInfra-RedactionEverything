"""Export/delivery job execution, format writers and the job quality-report zip."""
from __future__ import annotations

import csv
import json
import os
import sqlite3
import threading
import time
import uuid
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from app.core.config import settings
from app.services.structured_common import (
    MAX_EXPORT_ROWS,
    SUPPORTED_FILE_EXTENSIONS,
    normalize_value,
    quote_sqlite_ident,
    safe_filename,
    utc_iso,
)
from app.services.structured_files import iter_dataset_rows
from app.services.structured_masking import redact_row
from app.services.structured_profile import get_or_create_policy, profile_dataset
from app.services.structured_store import StructuredStore, get_structured_store

_STRUCTURED_EXPORT_LOCK = threading.Lock()


def export_base_filename(name: str) -> str:
    safe = safe_filename(name)
    suffix = Path(safe).suffix.lower()
    if suffix in SUPPORTED_FILE_EXTENSIONS:
        safe = safe[: -len(suffix)].strip("._")
    return safe or "data"


def unique_export_filename(
    filename: str,
    *,
    dataset_id: str,
    owner_id: str,
    job_id: str,
    out_dir: str,
    store: StructuredStore,
) -> str:
    """Keep job export names readable while preventing collisions in the delivery package."""
    existing_names = {
        str(export.get("filename") or "").lower()
        for export in store.list_exports(owner_id=owner_id, job_id=job_id)
        if export.get("filename")
    }
    try:
        existing_names.update(item.lower() for item in os.listdir(out_dir))
    except FileNotFoundError:
        pass

    if filename.lower() not in existing_names and not os.path.exists(os.path.join(out_dir, filename)):
        return filename

    stem, ext = os.path.splitext(filename)
    short_id = safe_filename(dataset_id).replace("-", "")[:8] or uuid.uuid4().hex[:8]
    for index in range(2, 1000):
        candidate = f"{stem}-{index}-{short_id}{ext}"
        if candidate.lower() not in existing_names and not os.path.exists(os.path.join(out_dir, candidate)):
            return candidate
    raise ValueError("unable to allocate a unique structured export filename")


class _ExportRowLimitExceeded(Exception):
    pass


def _write_export_parts(
    fmt: str,
    out_dir: str,
    base_filename: str,
    table_name: str,
    columns: list[str],
    rows_iter,
    rows_per_part: int | None,
) -> list[dict[str, Any]]:
    """流式写导出文件；xlsx 按 rows_per_part 分卷。失败时清掉已写文件再抛。"""
    stem, ext = os.path.splitext(base_filename)
    written_paths: list[str] = []
    parts: list[dict[str, Any]] = []

    def _part_path(index: int) -> tuple[str, str]:
        name = f"{stem}-part{index:03d}{ext}"
        return name, os.path.join(out_dir, name)

    try:
        if fmt == "csv":
            path = os.path.join(out_dir, base_filename)
            written_paths.append(path)
            count = 0
            with open(path, "w", encoding="utf-8-sig", newline="") as fh:
                writer = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
                writer.writeheader()
                for row in rows_iter:
                    writer.writerow(row)
                    count += 1
            parts.append({"filename": base_filename, "path": path, "rows": count})
        elif fmt == "xlsx":
            from openpyxl import Workbook

            cap = max(1, int(rows_per_part or 50_000))
            wb = ws = None
            part_rows = 0

            def _open(index: int):
                nonlocal wb, ws, part_rows
                name, path = _part_path(index)
                written_paths.append(path)
                parts.append({"filename": name, "path": path, "rows": 0})
                wb = Workbook(write_only=True)
                ws = wb.create_sheet("redacted")
                ws.append(columns)
                part_rows = 0

            def _save():
                nonlocal wb
                if wb is not None:
                    wb.save(parts[-1]["path"])
                    parts[-1]["rows"] = part_rows
                    wb = None

            _open(1)
            for row in rows_iter:
                if part_rows >= cap:
                    _save()
                    _open(len(parts) + 1)
                ws.append([row.get(col) for col in columns])
                part_rows += 1
            _save()
            if len(parts) == 1:
                # 单卷去 -part 后缀，保持既有交付命名
                plain_path = os.path.join(out_dir, base_filename)
                os.replace(parts[0]["path"], plain_path)
                parts[0]["filename"] = base_filename
                parts[0]["path"] = plain_path
        elif fmt == "sqlite":
            path = os.path.join(out_dir, base_filename)
            written_paths.append(path)
            if os.path.exists(path):
                os.remove(path)
            count = 0
            with sqlite3.connect(path) as conn:
                cols = ", ".join(f"{quote_sqlite_ident(col)} TEXT" for col in columns)
                conn.execute(f"CREATE TABLE {quote_sqlite_ident(table_name)} ({cols})")
                placeholders = ", ".join("?" for _ in columns)
                col_names = ", ".join(quote_sqlite_ident(col) for col in columns)
                batch: list[list[str]] = []
                for row in rows_iter:
                    batch.append([normalize_value(row.get(col)) for col in columns])
                    count += 1
                    if len(batch) >= 1000:
                        conn.executemany(
                            f"INSERT INTO {quote_sqlite_ident(table_name)} ({col_names}) VALUES ({placeholders})",
                            batch,
                        )
                        batch = []
                if batch:
                    conn.executemany(
                        f"INSERT INTO {quote_sqlite_ident(table_name)} ({col_names}) VALUES ({placeholders})",
                        batch,
                    )
                conn.commit()
            parts.append({"filename": base_filename, "path": path, "rows": count})
        elif fmt == "sql":
            path = os.path.join(out_dir, base_filename)
            written_paths.append(path)
            count = 0
            with open(path, "w", encoding="utf-8") as fh:
                cols = ", ".join(quote_sqlite_ident(col) for col in columns)
                fh.write(f"-- Redacted export generated at {utc_iso()}\n")
                for row in rows_iter:
                    values = ", ".join(sql_literal(row.get(col)) for col in columns)
                    fh.write(f"INSERT INTO {quote_sqlite_ident(table_name)} ({cols}) VALUES ({values});\n")
                    count += 1
            parts.append({"filename": base_filename, "path": path, "rows": count})
        else:
            raise ValueError(f"unsupported export format: {fmt}")
    except Exception:
        for path in written_paths:
            try:
                if os.path.exists(path):
                    os.remove(path)
            except OSError:
                pass
        raise
    return parts


def export_dataset(
    dataset_id: str,
    *,
    owner_id: str,
    job_id: str,
    export_format: str,
    store: StructuredStore | None = None,
) -> dict[str, Any]:
    store = store or get_structured_store()
    dataset = store.get_dataset(dataset_id, owner_id=owner_id)
    if not dataset:
        raise ValueError("dataset not found")
    policy = get_or_create_policy(dataset_id, owner_id=owner_id, store=store)
    policy_columns = policy.get("columns") or []
    fmt = "csv" if export_format == "zip" else export_format
    if fmt not in {"csv", "xlsx", "sqlite", "sql"}:
        raise ValueError(f"unsupported export format: {export_format}")
    out_dir = os.path.join(settings.OUTPUT_DIR, "structured", safe_filename(owner_id), job_id)
    os.makedirs(out_dir, exist_ok=True)
    base = export_base_filename(dataset.get("name") or dataset_id)

    columns, source_rows = iter_dataset_rows(dataset, owner_id=owner_id, store=store)
    max_rows = int(getattr(settings, "STRUCTURED_MAX_EXPORT_ROWS", 0) or MAX_EXPORT_ROWS)
    counter = {"rows": 0}

    def redacted_rows():
        for row in source_rows:
            counter["rows"] += 1
            if counter["rows"] > max_rows:
                raise _ExportRowLimitExceeded()
            yield redact_row(row, policy_columns, owner_id=owner_id, dataset_id=dataset_id)

    rows_per_part = int(getattr(settings, "EXPORT_TABLE_ROWS_PER_FILE", 50_000)) if fmt == "xlsx" else None

    with _STRUCTURED_EXPORT_LOCK:
        filename = unique_export_filename(
            f"{base}.{fmt}",
            dataset_id=dataset_id,
            owner_id=owner_id,
            job_id=job_id,
            out_dir=out_dir,
            store=store,
        )
        export_base = export_base_filename(filename)
        try:
            parts = _write_export_parts(
                fmt, out_dir, filename, export_base, columns, redacted_rows(), rows_per_part
            )
        except _ExportRowLimitExceeded:
            # 绝不静默截断：超限即失败，partial 已被 _write_export_parts 清理
            raise ValueError(
                f"数据集超过导出上限 {max_rows} 行（实际 ≥{counter['rows']}）；"
                "请分表导出或调高 STRUCTURED_MAX_EXPORT_ROWS"
            )

    total_rows = counter["rows"]
    action_counts = Counter(str(col.get("action") or "keep") for col in policy_columns if col.get("enabled", True))
    redacted_columns = sum(1 for col in policy_columns if col.get("enabled", True) and col.get("action") != "keep")
    base_summary = {
        "dataset_id": dataset_id,
        "dataset_name": dataset.get("name"),
        "row_count": total_rows,
        "column_count": len(columns),
        "redacted_column_count": redacted_columns,
        "action_counts": dict(action_counts),
        "shape_kind": dataset.get("shape_kind"),
        "part_count": len(parts),
        "total_rows": total_rows,
    }
    record: dict[str, Any] | None = None
    for index, part in enumerate(parts):
        record = store.add_export(
            owner_id=owner_id,
            job_id=job_id,
            dataset_id=dataset_id,
            export_format=fmt,
            file_path=part["path"],
            filename=part["filename"],
            summary={
                **base_summary,
                "export_filename": part["filename"],
                "part_index": index + 1,
                "part_rows": part["rows"],
            },
        )
    assert record is not None
    return record










def sql_literal(value: Any) -> str:
    if value is None:
        return "NULL"
    return "'" + normalize_value(value).replace("'", "''") + "'"


async def run_structured_job_item(
    *,
    job_id: str,
    item_id: str,
    dataset_id: str,
    owner_id: str,
    export_format: str,
    store: Any,
) -> dict[str, Any]:
    started = time.perf_counter()
    store.update_item_progress(
        item_id,
        stage="structured_profile",
        current=1,
        total=3,
        message="Profiling table columns",
    )
    profile = profile_dataset(dataset_id, owner_id=owner_id)
    store.update_item_progress(
        item_id,
        stage="structured_redaction",
        current=2,
        total=3,
        message="Applying column policy",
    )
    export = export_dataset(dataset_id, owner_id=owner_id, job_id=job_id, export_format=export_format)
    store.update_item_progress(
        item_id,
        stage="structured_export",
        current=3,
        total=3,
        message="Structured export ready",
    )
    return {
        "profile": {
            "sampled_rows": profile.get("sampled_rows"),
            "column_count": len(profile.get("columns") or []),
        },
        "export": export,
        "duration_ms": max(0, int((time.perf_counter() - started) * 1000)),
    }


def build_job_export_zip(*, owner_id: str, job_id: str, store: StructuredStore | None = None) -> str:
    store = store or get_structured_store()
    exports = store.list_exports(owner_id=owner_id, job_id=job_id)
    if not exports:
        raise ValueError("no structured exports ready")
    out_dir = os.path.join(settings.OUTPUT_DIR, "structured", safe_filename(owner_id), job_id)
    os.makedirs(out_dir, exist_ok=True)
    zip_path = os.path.join(out_dir, "structured-redacted.zip")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        manifest = {
            "job_id": job_id,
            "generated_at": utc_iso(),
            "exports": [
                {
                    "dataset_id": export.get("dataset_id"),
                    "filename": export.get("filename"),
                    "summary": export.get("summary") or {},
                }
                for export in exports
            ],
        }
        for export in exports:
            file_path = str(export.get("file_path") or "")
            if os.path.isfile(file_path):
                zf.write(file_path, arcname=str(export.get("filename") or os.path.basename(file_path)))
        zf.writestr("quality-report.json", json.dumps(manifest, ensure_ascii=False, indent=2))
    return zip_path


def structured_item_meta(dataset_id: str, *, owner_id: str, store: StructuredStore | None = None) -> dict[str, Any] | None:
    store = store or get_structured_store()
    dataset = store.get_dataset(dataset_id, owner_id=owner_id)
    if not dataset:
        return None
    policy = store.get_policy(dataset_id, owner_id=owner_id) or {}
    redacted_columns = sum(1 for col in policy.get("columns") or [] if col.get("enabled", True) and col.get("action") != "keep")
    return {
        "filename": dataset.get("name"),
        "file_type": f"structured:{dataset.get('source_kind')}",
        "has_output": False,
        "entity_count": redacted_columns,
        "metadata_warning": None,
    }
